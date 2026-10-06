from typing import List
import docx

from app.services.extractors.base import BaseExtractor, ExtractedPage


class DOCXExtractor(BaseExtractor):
    """Extracts text from DOCX files including paragraphs and tables."""

    def extract(self, file_path: str) -> List[ExtractedPage]:
        try:
            doc = docx.Document(file_path)
            lines: List[str] = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    lines.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        lines.append(row_text)

            full_text = "\n\n".join(lines).strip()
            if not full_text:
                return []

            return [
                ExtractedPage(
                    text=full_text,
                    page_number=1,
                    metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
                )
            ]
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
