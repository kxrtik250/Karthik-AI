from typing import List
from pptx import Presentation

from app.services.extractors.base import BaseExtractor, ExtractedPage


class PPTXExtractor(BaseExtractor):
    """Extracts text from PPTX presentations preserving slide numbers."""

    def extract(self, file_path: str) -> List[ExtractedPage]:
        pages: List[ExtractedPage] = []
        try:
            prs = Presentation(file_path)
            for idx, slide in enumerate(prs.slides, start=1):
                slide_texts: List[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                slide_texts.append(text)

                combined = "\n".join(slide_texts).strip()
                if combined:
                    pages.append(
                        ExtractedPage(
                            text=combined,
                            page_number=idx,
                            metadata={"slide": idx, "total_slides": len(prs.slides)},
                        )
                    )
        except Exception as e:
            raise ValueError(f"Failed to extract text from PPTX: {str(e)}")

        return pages
